"""Curate existing licensed, horizontal photographs and appropriate diagrams."""
import json,hashlib,copy,shutil
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
BASE=Path(__file__).resolve().parent

def read(p):return json.loads((BASE/p).read_text())
def save(p,x):(BASE/p).write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')

# Subjects, credits, and license links were checked on their Pexels pages.
sources={
'watch':(5349810,'different-devices-lying-on-table-near-notebook','Josh Sorenson','Đồng hồ và điện thoại đặt trên bàn, minh họa thiết bị theo dõi.','A watch and phone on a table, illustrating wearable devices.'),
'watch2':(32977239,'modern-fitness-tracking-smartwatch-and-smartphone','Andrey Matveev','Đồng hồ, nhẫn và điện thoại dùng để theo dõi sức khỏe; các số trên ảnh chỉ là ví dụ.','A watch, ring and phone; the displayed numbers are illustrative.'),
'bed':(4913754,'cozy-bed-with-blanket-and-soft-pillows','Julia Filirovska','Giường và gối trong phòng ngủ, minh họa môi trường nghỉ ngơi.','A bed and pillows illustrating a place to rest.'),
'bed2':(8251692,'photo-of-pillows-on-bed','PNW Production','Gối và chăn trên giường, minh họa không gian nghỉ ngơi.','Pillows and bedding illustrating a restful environment.'),
'gym':(11513151,'fitness-and-workout-equipment-on-the-floor','Frank Schrader','Tạ đòn và tạ tay trong không gian tập luyện.','A barbell and dumbbells in a workout space.'),
'dumbbells':(8032748,'close-up-photo-of-two-dumbbells-on-black-yoga-mat','MART PRODUCTION','Tạ tay đặt trên thảm tập, minh họa dụng cụ tập sức mạnh.','Dumbbells on a mat, illustrating strength-training equipment.'),
'workout':(8032758,'workout-equipment-on-black-yoga-mat','MART PRODUCTION','Dụng cụ vận động và chai nước đặt trên thảm tập.','Workout equipment and a water bottle on an exercise mat.'),
'ecg':(35044388,'cardiology-tools-with-ecg-and-stethoscope','Marta Branco','Dụng cụ y tế và bản ghi ECG; ảnh minh họa, không phải dữ liệu của người đọc.','Medical equipment and an ECG record; not the reader’s data.'),
'cups':(7368007,'clear-glass-teacup-on-the-table','Evgeniy Alekseyev','Tách trà đặt trên bàn, minh họa một khoảng thời gian dành cho trò chuyện.','Cups of tea illustrating time set aside for conversation.'),
'scale':(6975469,'measuring-tape-on-a-weighing-scale','SHVETS production','Cân và thước dây; cân nặng không tự cho biết tỷ lệ mỡ hay khối nạc.','A scale and measuring tape; weight alone does not show body composition.'),
'mat':(4793328,'rolled-yoga-mat-and-bottle-of-water-in-sunlight','Anete Lusina','Thảm tập và nước, minh họa bối cảnh vận động nhẹ.','An exercise mat and water, illustrating gentle activity.'),
'shoes':(135229,'pair-of-brown-low-top-sneakers-on-wooden-plank-floor','Steyn Viljoen','Đôi giày trên lối đi, minh họa hoạt động đi bộ.','Shoes by a walkway, illustrating walking.'),
'calendar':(28353112,'a-notebook-and-pen-on-a-white-table','Natalia S','Sổ lịch và bút, minh họa việc theo dõi thời điểm và triệu chứng.','A planner and pen, illustrating tracking dates and symptoms.'),
'journal':(9429448,'opened-notebook-with-white-blank-pages-and-pen-on-table','Monstera Production','Sổ mở và bút, minh họa cách ghi lại câu hỏi hoặc diễn biến.','An open notebook and pen, illustrating notes or questions.'),
}
pool={}
for key,(num,slug,author,vi,en) in sources.items():
 p=BASE/f'images/researched/landscape-2026-10-29/{num}.jpg';im=Image.open(p);assert im.width>im.height
 url=f'https://www.pexels.com/photo/{slug}-{num}/'
 pool[key]=dict(asset=str(p.relative_to(BASE)),title=slug.replace('-',' '),source_page=url,source_asset_url=f'https://images.pexels.com/photos/{num}/pexels-photo-{num}.jpeg',download_url=f'https://images.pexels.com/photos/{num}/pexels-photo-{num}.jpeg?auto=compress&cs=tinysrgb&w=1600&fm=jpg',author=author,license='Pexels License',license_url='https://www.pexels.com/license/',caption_vi=vi,caption_en=en,alt_vi=vi,alt_en=en,fit='cover',modifications_vi='Bản JPEG thu nhỏ/nén từ nguồn.',modifications_en='Resized/compressed JPEG from the source.',width=im.width,height=im.height,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),checked_date='2026-10-29',review_status='Subject, source and licence checked',source_metadata={'source_page':url,'license_page':'https://www.pexels.com/license/','photographer':author},relevance_vi='Ảnh nguồn thể hiện bối cảnh chủ đề; không sử dụng làm kết quả đo hay chẩn đoán.')
# Retain already-reviewed horizontal, non-person photographs where they are more specific.
existing=read('data/image_selections.json')['assets']
for key,aid,role in [('sauna','DM_HW_02','hero'),('hotwater','DM_WH_03','hero'),('meal','MW_NUT_01','hero'),('salad','MW_NUT_02','hero'),('prep','MW_NUT_07','hero'),('breakfast','WH_21','hero'),('alcohol','MW_NUT_04','hero'),('blood','WH_22','inline'),('sperm','MW_NUT_06','inline')]:
 a=copy.deepcopy(next(a for a in existing if a['article_id']==aid and a['role']==role));im=Image.open(BASE/a['asset']);assert im.width>im.height,(key,im.size);a['fit']='cover';pool[key]=a
# Each article is mapped to an actual subject, not merely an available image.
groups={
'watch':'LM_HW_01 LM_RC_02 LM_AC_01 LM_AC_02 LM_HW_02 WH_16 WHOOP-ART-02 WHOOP-ART-04 WHOOP-ART-05 WHOOP-ART-09 WHOOP-ART-20 WHOOP-ART-25',
'watch2':'LM_RC_01 WHOOP-ART-03 WHOOP-ART-07',
'bed':'LM_SL_01 LM_SL_02 DM_SL_01 DM_SL_02 WH_19 MW_SXH_03 MW_REL_08 WHOOP-ART-13 WHOOP-ART-15 WHOOP-ART-16 WHOOP-ART-21',
'bed2':'WHOOP-ART-14 WHOOP-ART-24 MW_SXH_05',
'gym':'DM_AC_01 DM_AC_02 WH_17 WH_18 WHOOP-ART-01 WHOOP-ART-26',
'dumbbells':'MW_NUT_03 MW_SXH_06',
'workout':'DM_RC_02 DM_WH_01 WH_15 WHOOP-ART-08',
'mat':'DM_RC_01 DM_HW_01 DM_MH_02 WH_26 WH_27 WHOOP-ART-10 WHOOP-ART-11 WHOOP-ART-23',
'shoes':'WHOOP-ART-27',
'ecg':'MW_SXH_02 MW_SXH_08 WHOOP-ART-12 WHOOP-ART-28 WHOOP-ART-29',
'calendar':'LM_WH_01 LM_WH_02 LM_WH_03 DM_WH_02 WH_04 WH_06 WH_07 WH_10 WH_16 WH_20 WH_23 WH_24 WH_28',
'journal':'MW_SXO_05 LM_MH_01 LM_MH_02 LM_MH_03 DM_MH_01 DM_MH_03 WH_09 WH_11 WH_12 WH_13 WH_25 WH_29 WH_30 MW_SXH_01 MW_SXH_04 MW_SXO_01 MW_SXO_03 MW_SXO_04 MW_SXO_06 MW_SXO_08 WHOOP-ART-22',
'cups':'MW_REL_01 MW_REL_02 MW_REL_03 MW_REL_04 MW_REL_05 MW_REL_06 MW_REL_07 MW_SXO_02 WHOOP-ART-22',
'scale':'WHOOP-ART-17 WHOOP-ART-18',
'sauna':'DM_HW_02',
'hotwater':'DM_WH_03',
'meal':'MW_NUT_01 WHOOP-ART-06',
'salad':'MW_NUT_02 MW_NUT_05',
'prep':'MW_NUT_07',
'breakfast':'WH_21',
'alcohol':'MW_NUT_04',
'blood':'WH_22',
'sperm':'MW_NUT_06 MW_SXH_07',
}
# Specific topics supersede a broad category. These keep the subject closely aligned.
overrides={'WHOOP-ART-08':'bed2','LM_RC_02':'watch2','WH_16':'watch2','WHOOP-ART-03':'shoes','WHOOP-ART-04':'watch2','WHOOP-ART-07':'meal','WHOOP-ART-10':'journal','WHOOP-ART-11':'journal','WHOOP-ART-12':'ecg','WHOOP-ART-19':'meal','WHOOP-ART-22':'journal','WH_26':'mat','WH_27':'mat','LM_MH_01':'journal','LM_MH_02':'mat','LM_MH_03':'journal','DM_MH_01':'journal','DM_MH_02':'mat','WH_13':'journal','MW_NUT_05':'salad','MW_NUT_06':'sperm','MW_NUT_07':'prep','MW_NUT_08':'workout','MW_SXH_07':'sperm','WH_22':'blood','WH_21':'breakfast','WH_29':'calendar','WH_30':'journal','WHOOP-ART-27':'shoes'}
assign={}
for key,ids in groups.items():
 for aid in ids.split():assign[aid]=key
assign.update(overrides)
articles=read('articles.json');missing=[a['id'] for a in articles if a['id'] not in assign];assert not missing,missing
# Only retain source-backed, horizontal, non-person figures that explain the associated topic.
keep={'LM_SL_02','LM_RC_01','DM_RC_01','LM_WH_01','WH_11','WH_12','WH_13','WH_16','WH_22','WH_27','MW_SXH_02','MW_SXO_02','MW_NUT_06','MW_SXH_07'}
oldsel=read('data/image_selections.json');newsel=[];custom={}
for a in articles:
 aid=a['id'];item=copy.deepcopy(pool[assign[aid]]);item.update(article_id=aid,role='hero',batch=a.get('image_review_batch',1) or 1)
 # Remove metadata not part of the image record schema.
 hero={k:v for k,v in item.items() if k!='source_metadata'}
 if aid.startswith('WHOOP-ART-'):
  custom[str(int(aid[-2:]))]=hero
  continue
 old_inline=next((x for x in oldsel['assets'] if x['article_id']==aid and x['role']=='inline'),None)
 entries=[item]
 if aid in keep:
  assert old_inline is not None,aid
  im=Image.open(BASE/old_inline['asset']);assert im.width>im.height,(aid,im.size)
  entries.append(old_inline)
 newsel.extend(entries)
 a.update(hero_image=hero['asset'],thumbnail=hero['asset'],hero_image_fit='cover',hero_image_source_url=hero['source_page'],hero_caption=hero['caption_en'],hero_caption_vi=hero['caption_vi'],image_attributions=[{k:v for k,v in e.items() if k!='source_metadata'} for e in entries],image_provenance='Ảnh khổ ngang đã đối chiếu nội dung, tác giả và giấy phép; hình chuyên môn giữ nguồn.',image_review_status='Landscape image subject, source and licence checked')
 if aid not in keep:a.update(inline_image='',inline_image_source_url='',inline_caption='',inline_caption_vi='')
oldsel['assets']=newsel;save('data/image_selections.json',oldsel);save('data/whoop_custom_heroes.json',custom);save('data/landscape_sources.json',pool)
save('articles.json',articles)
# Keep the merged image manifest and portable long-form snapshot coherent.
manifest=read('images/image_manifest.json');longform=read('data/longform_articles.json')
for a in articles:
 aid=a['id']
 if aid.startswith('WHOOP-ART-'):continue
 m=manifest[aid];m.update(hero_local=a['hero_image'],figure_local=a['inline_image'],hero_source_page=a['hero_image_source_url'],figure_source_page=a['inline_image_source_url'],image_attributions=a['image_attributions'],provenance=a['image_provenance'])
 for k in ['hero_image','thumbnail','hero_image_source_url','inline_image','inline_image_source_url','inline_caption','inline_caption_vi','hero_caption','hero_caption_vi','hero_image_fit','image_attributions','image_provenance','image_review_status']:
  longform[aid][k]=a[k]
save('images/image_manifest.json',manifest);save('data/longform_articles.json',longform)
# Update only the affected image fields in the existing workbook.
book=BASE/'docs/Master_Health_Articles_MVP.xlsx';wb=load_workbook(book);ws=wb['Master_Content_MVP'];header={c.value:c.column for c in ws[1]};byid={a['id']:a for a in articles}
fields={'Thumbnail_Asset':'thumbnail','Hero_Image_Asset':'hero_image','Hero_Image_Source_URL':'hero_image_source_url','Inline_Figure_Asset':'inline_image','Inline_Figure_Caption_VI':'inline_caption_vi','Inline_Figure_Caption_EN':'inline_caption','Inline_Figure_Source_URL':'inline_image_source_url','Hero_Image_Caption_VI':'hero_caption_vi','Hero_Image_Caption_EN':'hero_caption','Image_Attributions_JSON':'image_attributions','Image_Provenance':'image_provenance','Image_Review_Status':'image_review_status'}
for row in range(2,ws.max_row+1):
 a=byid[ws.cell(row,header['Content_ID']).value]
 for col,key in fields.items():
  val=a[key];val=json.dumps(val,ensure_ascii=False) if isinstance(val,list) else val;c=ws.cell(row,header[col],val);c.hyperlink=None
  if col in ['Hero_Image_Source_URL','Inline_Figure_Source_URL'] and val:c.hyperlink=val
reg=wb['Image_Register'];reg.delete_rows(2,reg.max_row-1);columns=[c.value for c in reg[1]]
for item in newsel:
 reg.append([item.get(k,'') for k in columns])
 for name in ['source_page','license_url']:
  cell=reg.cell(reg.max_row,columns.index(name)+1)
  if cell.value:cell.hyperlink=cell.value
wb.save(book);other=BASE.parent/'Fitness IOS'/book.name
if other.exists():shutil.copy2(book,other)
print('Mapped',len(assign),'articles; selected',len(pool),'landscape non-person sources; retained',len(keep),'relevant figures.')
