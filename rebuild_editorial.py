"""Build both languages, first-party visuals, Markdown and Excel from the reviewed TSVs.

Run from any directory with Python 3, openpyxl and markdown installed.
The TSV files are the editorial source of truth. Legacy batch builders are historical.
"""
import argparse
import csv
from collections import Counter
from copy import copy
from datetime import date
import html
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import sys
import textwrap

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
AUDIT = ROOT / 'audit' / 'health-2026-10-26'
FITNESS = ROOT / 'Fitness IOS'
REVIEW_DATE = '2026-10-26'
LIVE_URL = 'https://test-repo-y1wb.onrender.com'

CONFIRMED_MISMATCHES = {
    'LM_SL_01', 'LM_RC_02', 'LM_HW_01', 'LM_HW_02', 'LM_WH_01', 'LM_WH_02',
    'LM_WH_03', 'DM_WH_01', 'LM_MH_01', 'LM_MH_03', 'DM_MH_01', 'DM_MH_03',
    'WH_06', 'WH_07', 'WH_09', 'WH_10', 'MW_NUT_02', 'MW_NUT_04', 'MW_SXH_04',
    'MW_REL_03', 'MW_REL_04', 'MW_SXO_01', 'MW_SXO_04', 'MW_SXH_05',
    'MW_SXH_06', 'MW_SXH_07', 'MW_REL_07', 'MW_REL_08', 'MW_SXO_05',
}
FIGURE_EXAMPLES = {
    'LM_AC_02': 'Sơ đồ CO2/đáp ứng dịch trong hồi sức không giải thích Day Strain.',
    'WH_26': 'Hình thủ thuật sản khoa không hướng dẫn vận động thai kỳ.',
    'MW_REL_01': 'Sơ đồ giải phẫu thần kinh không minh chứng tỷ lệ giao tiếp 5:1.',
    'DM_SL_02': 'Công thức phân tử melatonin không giải thích tích lũy nợ ngủ.',
}

def read_tsv(name):
    with (BASE / 'data' / name).open(encoding='utf-8', newline='') as handle:
        rows = list(csv.DictReader(handle, delimiter='\t'))
    assert all(None not in r and all(v is not None for v in r.values()) for r in rows), name
    return rows

def save_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

def first_sentence(value):
    # Decimal values do not end a sentence.
    return re.split(r'(?<=[.!?])\s+(?=[A-ZÀ-ỸĐ])', value)[0]

def svg_text(value, x, y, width, size=25, fill='#183c3b', line_height=None):
    line_height = line_height or size * 1.42
    lines = textwrap.wrap(value, width=width, break_long_words=False, break_on_hyphens=False)
    return '\n'.join(
        f'<text x="{x}" y="{y + i * line_height:.1f}" font-size="{size}" fill="{fill}">{html.escape(line)}</text>'
        for i, line in enumerate(lines)
    ), len(lines) * line_height

def make_svg(row, category, kind):
    palette = {
        'Sleep': ('#293d72', '#ebeffa'), 'Recovery': ('#126453', '#e8f5ef'),
        'Activity': ('#8a4c14', '#fff3e3'), 'Health & Wellness': ('#145b75', '#e7f4fa'),
        "Women's Health": ('#843d67', '#fbeef6'), "Men's Health": ('#364d83', '#eef1fb'),
    }
    ink, bg = palette[category]
    if kind == 'hero':
        vi, _ = svg_text(row['title_vi'], 76, 205, 39, size=44, fill=ink)
        en, _ = svg_text(row['title_en'], 80, 430, 55, size=29, fill=ink)
        # Abstract editorial art. No invented measurements or anatomical claims.
        art = f'''<circle cx="1090" cy="350" r="138" fill="{ink}" opacity=".10"/>
          <rect x="970" y="250" width="220" height="230" rx="25" fill="white" stroke="{ink}" stroke-width="3"/>
          <path d="M1005 300h145m-145 35h110m-110 35h125" stroke="{ink}" opacity=".4" stroke-width="8" stroke-linecap="round"/>
          <path d="M1020 420l28 28 70-70" fill="none" stroke="{ink}" stroke-width="10" stroke-linecap="round" stroke-linejoin="round"/>'''
        body = f'''<rect width="1280" height="720" fill="{bg}"/>
          <text x="80" y="90" font-size="22" fill="{ink}" letter-spacing="3">PULSE &amp; VITALITY / {row['id']}</text>
          {vi}{en}{art}
          <path d="M80 595h1120" stroke="{ink}" opacity=".2"/>
          <text x="80" y="642" font-size="20" fill="{ink}">Hiểu bằng chứng • Chọn hành động phù hợp</text>
          <text x="80" y="677" font-size="17" fill="{ink}">Editorial illustration / Minh họa biên tập</text>'''
        height = 720
    else:
        body = f'<rect width="1280" height="1200" fill="{bg}"/>'
        title, _ = svg_text(row['title_vi'], 60, 76, 65, size=29, fill=ink)
        body += title
        body += f'<text x="62" y="164" font-size="19" fill="{ink}">TÓM TẮT / SUMMARY • {row["id"]} • {REVIEW_DATE}</text>'
        for i, (key, heading_vi, heading_en) in enumerate([
            ('facts', 'Hiểu đúng', 'Understand'), ('action', 'Thực hành', 'Try'), ('limits', 'Giới hạn', 'Limits')
        ]):
            y = 210 + i * 290
            body += f'<rect x="50" y="{y}" width="1180" height="270" rx="22" fill="white"/>'
            body += f'<text x="78" y="{y+42}" font-size="23" font-weight="700" fill="{ink}">{i+1:02d} / {heading_vi} · {heading_en}</text>'
            vi, vh = svg_text(first_sentence(row[key+'_vi']), 80, y+88, 47, size=21)
            en, eh = svg_text(first_sentence(row[key+'_en']), 675, y+88, 47, size=21)
            assert max(vh, eh) <= 175, (row['id'], key, vh, eh)
            body += vi + en
            body += f'<path d="M640 {y+66}v175" stroke="{ink}" opacity=".15"/>'
        body += f'<text x="65" y="1130" font-size="19" fill="{ink}">Minh họa biên tập; không phải dữ liệu đo hoặc hình nghiên cứu gốc.</text>'
        body += f'<text x="65" y="1165" font-size="18" fill="{ink}">Editorial summary; not measured data or a figure from a cited study. Read the full article for context.</text>'
        height = 1200
    return f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 {height}" width="1280" height="{height}" role="img" aria-labelledby="title desc">
      <title id="title">{html.escape(row['title_vi'])} / {html.escape(row['title_en'])}</title>
      <desc id="desc">Original editorial illustration. Not clinical data, an anatomical diagram or a reproduction of a study.</desc>
      <g font-family="Arial, Helvetica, sans-serif">{body}</g></svg>\n'''

def body_markdown(row, refs, lang):
    headings = ['Hiểu đúng', 'Bạn có thể làm gì?', 'Giới hạn cần biết', 'Khi nào cần tư vấn?'] if lang == 'vi' else [
        'Understanding the topic', 'What you can do', 'What the evidence cannot tell you', 'When to seek advice']
    parts = []
    # Direct source links accompany the central claim, with full support scopes below.
    for field, heading in zip(['facts', 'action', 'limits', 'care'], headings):
        value = row[field+'_'+lang]
        if field == 'facts':
            value += ' ' + ' '.join(f'[{i}]({ref["url"]})' for i,ref in enumerate(refs,1))
        parts += [f'## {heading}', value]
    parts += ['## Nguồn và phạm vi hỗ trợ' if lang == 'vi' else '## Sources and what they support']
    for n, ref in enumerate(refs, 1):
        parts += [f'{n}. [{ref["title"]}]({ref["url"]}) — {ref["kind"]}. '
                  + ('Hỗ trợ: ' if lang == 'vi' else 'Supports: ') + ref['support_'+lang] + '.']
    parts += [f'> Cập nhật nội dung và nguồn: {REVIEW_DATE}. Nội dung giáo dục; chưa có xác nhận duyệt bởi bác sĩ.'
              if lang == 'vi' else f'> Content and sources updated: {REVIEW_DATE}. Educational content; clinical sign-off has not been recorded.']
    return '\n\n'.join(parts) + '\n'

def table_sheet(wb, name, headers, rows):
    if name in wb:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor='184E52')
        c.font = Font(color='FFFFFF', bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for n, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(n)].width = 24 if n < 3 else 55
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 86
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(c.value, str) and c.value.startswith('https://') and '\n' not in c.value:
                c.hyperlink = c.value
                c.font = Font(color='0563C1', underline='single')
    return ws

def sync_master(articles, audit_rows, sources, link_results):
    source_path = FITNESS / 'Master_Health_Articles_MVP.xlsx'
    if not source_path.exists():
        source_path = BASE / 'docs' / source_path.name
    wb = openpyxl.load_workbook(source_path)
    fields = {
        'Content_ID':'id', 'Content_Type':'type', 'Category':'category', 'Sub_Topic':'sub_topic',
        'Target_Gender':'gender', 'Mapped_Screen':'screen', 'Mapped_Metric':'metric', 'Trigger_Condition':'trigger',
        'Title_VI':'title_vi', 'Short_Description_VI':'description_vi', 'Key_Takeaway_VI':'takeaway_vi',
        'Reading_Time_VI':'reading_time_vi', 'Full_Content_Markdown_VI':'content_vi',
        'Title_EN':'title', 'Short_Description':'description', 'Key_Takeaway':'takeaway',
        'Reading_Time':'reading_time', 'Full_Content_Markdown':'content',
        'Thumbnail_Asset':'thumbnail', 'Hero_Image_Asset':'hero_image', 'Hero_Image_Source_URL':'hero_image_source_url',
        'Inline_Figure_Asset':'inline_image', 'Inline_Figure_Caption_VI':'inline_caption_vi',
        'Inline_Figure_Caption_EN':'inline_caption', 'Inline_Figure_Source_URL':'inline_image_source_url',
        'Source_App_Name':'source_name', 'Source_URL':'source_url',
    }
    extra = {
        'Reference_List_JSON':'references', 'Evidence_Limitations_VI':'evidence_limitations_vi',
        'Evidence_Limitations_EN':'evidence_limitations_en', 'Source_Checked_Date':'source_checked_date',
        'Editorial_Review_Status':'editorial_review_status', 'Clinical_Review_Status':'clinical_review_status',
        'Image_Provenance':'image_provenance', 'Public_Article_URL':'public_url',
    }
    by_id = {a['id']:a for a in articles}
    ws = wb['Master_Content_MVP']
    header = {c.value:c.column for c in ws[1]}
    for key in extra:
        if key not in header:
            header[key] = ws.max_column + 1
            cell = ws.cell(1, header[key], key)
            cell._style = copy(ws.cell(1, 28)._style)
            ws.column_dimensions[cell.column_letter].width = 50
    for row in range(2, ws.max_row+1):
        art = by_id[ws.cell(row, header['Content_ID']).value]
        for column, key in {**fields, **extra}.items():
            val = art[key]
            if isinstance(val, list):
                val = json.dumps(val, ensure_ascii=False)
            cell = ws.cell(row, header[column], val)
            # Old Excel hyperlinks must not survive after replacing their display text.
            cell.hyperlink = None
            if column in ['Source_URL', 'Hero_Image_Source_URL', 'Inline_Figure_Source_URL', 'Public_Article_URL']:
                cell.hyperlink = val
        ws.cell(row, header['Status'], 'Revised — clinical sign-off not recorded')
    schema=wb['HDSD_Data_Schema']
    example=articles[0]
    schema_changes={
        'Hero_Image_Source_URL':('Original editorial artwork','Mở minh họa gốc do dự án tạo','URL của SVG biên tập; không phải ảnh Unsplash hay bằng chứng y khoa.'),
        'Inline_Figure_Source_URL':('Original editorial summary','Mở sơ đồ tóm tắt do dự án tạo','URL của SVG biên tập; không phải hình sao chép từ nghiên cứu.'),
        'Source_URL':('Primary reference','Nguồn tham khảo chính','Một nguồn chính; xem Reference_List_JSON và Source_Register để biết đầy đủ nguồn và phạm vi hỗ trợ.'),
        'Source_App_Name':('Reference label','Tên nguồn tham khảo chính','Tên cơ quan hoặc nghiên cứu; không hàm ý cơ quan đó bảo trợ website.'),
    }
    example_fields=['Source_URL','Source_App_Name','Hero_Image_Source_URL','Inline_Figure_Source_URL',
                    'Thumbnail_Asset','Hero_Image_Asset','Inline_Figure_Asset','Inline_Figure_Caption_VI','Inline_Figure_Caption_EN']
    existing_schema={schema.cell(r,1).value:r for r in range(2,schema.max_row+1)}
    for r in range(2,schema.max_row+1):
        key=schema.cell(r,1).value
        if key in schema_changes:
            for col,val in enumerate(schema_changes[key],2):schema.cell(r,col,val)
        if key in example_fields:
            c=schema.cell(r,5,example[fields[key]]);c.hyperlink=None
            if str(c.value).startswith('https://'):c.hyperlink=c.value
        if key=='Status':schema.cell(r,5,'Revised — clinical sign-off not recorded')
    for key,value in extra.items():
        vals=[key,'Editorial metadata','Chi tiết nguồn và trạng thái rà soát','Phân biệt đối chiếu nguồn với xác nhận duyệt bởi bác sĩ.',
              json.dumps(example[value],ensure_ascii=False) if isinstance(example[value],list) else example[value]]
        r=existing_schema.get(key,schema.max_row+1)
        for col,val in enumerate(vals,1):schema.cell(r,col,val)
    mm = wb['Mapping_Matrix']
    mh = {c.value:c.column for c in mm[1]}
    for r in range(2, mm.max_row+1):
        a = by_id[mm.cell(r,mh['Content_ID']).value]
        for col,key in [('Title_VI','title_vi'),('Title_EN','title'),('Source_App_Name','source_name'),('Canonical_Source_URL','source_url')]:
            c=mm.cell(r,mh[col],a[key]);c.hyperlink=None
            if col=='Canonical_Source_URL':c.hyperlink=a[key]
        mm.cell(r,mh['Mapped_Metric'],a['metric'])
        mm.cell(r,mh['Sub_Topic'],a['sub_topic'])
    table_sheet(wb, 'Editorial_Audit', list(audit_rows[0]), [list(r.values()) for r in audit_rows])
    table_sheet(wb, 'Source_Register', ['Source_ID','Title','URL','Source_Type','Supports_VI','Supports_EN','Checked_Date','Verification_Note'], [
        [s['id'],s['title'],s['url'],s['kind'],s['support_vi'],s['support_en'],REVIEW_DATE,
         'Content reviewed using primary publisher / indexed record; HTTP availability logged separately.'] for s in sources.values()])
    table_sheet(wb, 'Link_Audit', ['Original_URL','HTTP_Status','Final_URL','Page_Title','Interpretation'], [
        [u,r.get('http_status','NETWORK_ERROR'),r.get('final_url',''),r.get('page_title',''),
         'Transport only. 403 / 203 / challenge pages are not proof a source is invalid; title and claim support require separate review.']
        for u,r in sorted(link_results.items())])
    source_path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(source_path)
    web_copy=BASE/'docs'/source_path.name
    web_copy.parent.mkdir(parents=True,exist_ok=True)
    if source_path.resolve()!=web_copy.resolve():shutil.copy2(source_path,web_copy)

def sync_roadmaps(articles):
    by_id={a['id']:a for a in articles}
    for name in ['Mens_Wellness_Content_Roadmap.xlsx', 'Master_Health_Articles_MVP - 30 Blog Topics.xlsx']:
        original=FITNESS / name
        if not original.exists():original=BASE / 'docs' / name
        if not original.exists():continue
        wb=openpyxl.load_workbook(original)
        if name.startswith('Mens'):
            ws=wb['Content Roadmap'];h={c.value:c.column for c in ws[1]}
            for r in range(2,ws.max_row+1):
                aid=ws.cell(r,h['Assigned_Content_ID']).value
                if aid not in by_id:
                    ws.cell(r,h['Content status'],'No matching article — review mapping')
                    ws.cell(r,h['Evidence status'],'Not assessed — no matching article')
                    continue
                a=by_id[aid]
                for col,key in [('Article title (EN)','title'),('Tiêu đề (VI)','title_vi'),('Nội dung chính','description_vi')]:
                    ws.cell(r,h[col],a[key])
                for col,key in [('Scientific source / research URL','source_url'),('Elevated_Title_VI','title_vi'),('Elevated_Title_EN','title'),('Key_Takeaways_VI','takeaway_vi'),('Biometric_Formula_Mapping','metric')]:
                    c=ws.cell(r,h[col],a[key]);c.hyperlink=None
                    if key=='source_url':c.hyperlink=a[key]
                ws.cell(r,h['Content status'],f'Revised — {aid}; clinical sign-off not recorded')
                ws.cell(r,h['Evidence status'],'Source replaced / checked; see master Editorial_Audit')
        else:
            ws=wb.worksheets[0]
            for r in range(5,ws.max_row+1):
                aid=ws.cell(r,6).value
                if aid not in by_id:continue
                a=by_id[aid]
                for col,key in [(3,'title'),(7,'title'),(8,'title_vi'),(9,'description_vi'),(10,'takeaway_vi'),(11,'metric')]:
                    ws.cell(r,col,a[key])
                ws.cell(r,5,f'Revised — {aid}; clinical sign-off not recorded')
                c=ws.cell(r,12,a['source_url'])
                c.hyperlink=a['source_url']
        # A linked, complete table is safer than assuming the display layout's column positions.
        subset=[a for a in articles if a['id'].startswith('MW_') or a['category']=="Men's Health"] if name.startswith('Mens') else [a for a in articles if a['category']=="Women's Health"]
        table_sheet(wb,'Reviewed_Content', ['Content_ID','Title_VI','Title_EN','Source_Name','Source_URL','Public_URL','Review_Status','Limits_VI'], [
            [a['id'],a['title_vi'],a['title'],a['source_name'],a['source_url'],a['public_url'],a['editorial_review_status']+'; clinical sign-off not recorded',a['evidence_limitations_vi']] for a in subset])
        all_refs={ref['id']:ref for a in subset for ref in a['references']}
        table_sheet(wb,'Source_Register',['Source_ID','Title','URL','Type','Supports_VI','Supports_EN'],[
            [r['id'],r['title'],r['url'],r['kind'],r['support_vi'],r['support_en']] for r in all_refs.values()])
        wb.save(original)
        web_copy=BASE/'docs'/name
        if original.resolve()!=web_copy.resolve():shutil.copy2(original,web_copy)

def main():
    sources={s['id']:s for s in read_tsv('editorial_sources.tsv')}
    rows=[r for name in ['editorial_articles.tsv','editorial_women.tsv','editorial_men.tsv'] for r in read_tsv(name)]
    catalog={r['id']:r for r in rows}
    articles=json.loads((BASE/'articles.json').read_text())
    assert len(catalog)==len(rows)==len(articles)==83
    assert set(catalog)=={a['id'] for a in articles}
    AUDIT.mkdir(parents=True,exist_ok=True)
    backup=AUDIT/'before'
    backup.mkdir(exist_ok=True)
    for p in [BASE/'articles.json',BASE/'images/image_manifest.json']+list(FITNESS.glob('*.xlsx')):
        dst=backup/p.name
        if not dst.exists():shutil.copy2(p,dst)
    before=json.loads((backup/'articles.json').read_text())
    old={a['id']:a for a in before}
    metadata_path=AUDIT/'original_pubmed_metadata.json'
    metadata=json.loads(metadata_path.read_text())['resultList']['result'] if metadata_path.exists() else []
    pubmed={m['id']:m for m in metadata}
    manifest={};audit_rows=[]
    for a in articles:
        row=catalog[a['id']]
        # Old app mappings claimed inference of hormones, implantation and molecular
        # pathways from wearables. Preserve the association with the existing screen,
        # while replacing that unsupported interpretation with the actual topic.
        a['metric']=row['title_en']+' (educational topic; not a measured biomarker)'
        a['sub_topic']=row['title_en']
        refs=[dict(sources[k],checked_date=REVIEW_DATE) for k in row['sources'].split(',')]
        a.update(title=row['title_en'],title_vi=row['title_vi'],description=first_sentence(row['facts_en']),
                 description_vi=first_sentence(row['facts_vi']),takeaway=row['action_en'],takeaway_vi=row['action_vi'],
                 references=refs,source_name=refs[0]['title'],source_url=refs[0]['url'],
                 content=body_markdown(row,refs,'en'),content_vi=body_markdown(row,refs,'vi'),
                 evidence_limitations_vi=row['limits_vi'],evidence_limitations_en=row['limits_en'],
                 source_checked_date=REVIEW_DATE,editorial_review_status='Revised against cited sources',
                 clinical_review_status='Not recorded',image_provenance='Original project-owned SVG editorial illustration; not a study figure',
                 public_url=f'{LIVE_URL}/posts/{a["id"]}.html',
                 inline_caption_vi='Sơ đồ tóm tắt do biên tập tạo. Không phải số liệu đo, hình giải phẫu hay hình lấy từ nghiên cứu trích dẫn.',
                 inline_caption='Original editorial summary. Not measured data, an anatomical diagram or a figure reproduced from a cited study.')
        for lang,body_key,time_key in [('vi','content_vi','reading_time_vi'),('en','content','reading_time')]:
            # Exclude the bibliography from reading time.
            words=len(' '.join(row[k+'_'+lang] for k in ['facts','action','limits','care']).split())
            mins=max(1,math.ceil(words/180))
            a[time_key]=f'{mins} phút' if lang=='vi' else f'{mins} min'
        for kind,field in [('hero','hero_image'),('summary','inline_image')]:
            relative=f'images/editorial/{a["id"]}_{kind}.svg'
            p=BASE/relative;p.parent.mkdir(parents=True,exist_ok=True)
            p.write_text(make_svg(row,a['category'],kind),encoding='utf-8')
            a[field]=relative
            a['hero_image_source_url' if kind=='hero' else 'inline_image_source_url']=f'{LIVE_URL}/{relative}'
            if FITNESS.exists():
                dest=FITNESS/relative;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
        a['thumbnail']=a['hero_image']
        manifest[a['id']]={
            'hero_local':a['hero_image'],'figure_local':a['inline_image'],
            'hero_source_page':a['hero_image_source_url'],'figure_source_page':a['inline_image_source_url'],
            'creator':'Pulse & Vitality editorial assets','provenance':a['image_provenance'],
            'license_note':'Original SVG assets created in this project; no external photo or figure attribution claimed.',
            'evidence_urls':[r['url'] for r in refs],
        }
        md=f'# {a["title_vi"]}\n\n## English title\n\n{a["title"]}\n\n---\n\n'+a['content_vi']+'\n---\n\n'+a['content']
        for folder in [BASE/'markdown',FITNESS/'markdown']:
            if folder.parent.exists():folder.mkdir(exist_ok=True);(folder/f'{a["id"]}.md').write_text(md,encoding='utf-8')
        previous=old[a['id']];match=re.search(r'pubmed.ncbi.nlm.nih.gov/(\d+)',previous['source_url'])
        actual=pubmed.get(match[1],{}).get('title','') if match else ''
        audit_rows.append({
            'Content_ID':a['id'],'Title_VI':a['title_vi'],'Original_Source_URL':previous['source_url'],
            'Original_PMID_Actual_Title':actual,
            'Finding':'Confirmed unrelated canonical PMID' if a['id'] in CONFIRMED_MISMATCHES else 'Original body and citations replaced with bounded, directly cited guidance; no claim that every original reference was invalid',
            'Content_Fix_VI':row['limits_vi'],
            'Image_Finding_VI':FIGURE_EXAMPLES.get(a['id'],'Không giữ nhãn ảnh y khoa xác thực khi chưa có đủ đối chiếu nội dung và quyền sử dụng.'),
            'Image_Fix':'Replaced hero and inline figure with original, explicitly labelled editorial SVGs',
            'Revised_Source_URL':a['source_url'],'All_Source_URLs':'\n'.join(r['url'] for r in refs),
            'Action_VI':row['action_vi'],'Care_VI':row['care_vi'],
            'Checked_Date':REVIEW_DATE,'Clinical_Review':'Not recorded','Public_URL':a['public_url'],
        })
    save_json(BASE/'articles.json',articles)
    (BASE/'articles_data.js').write_text('const articles = '+json.dumps(articles,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')
    save_json(BASE/'images/image_manifest.json',manifest)
    save_json(BASE/'data/references.json',list(sources.values()))
    save_json(AUDIT/'article_audit.json',audit_rows)
    original_results_path=AUDIT/'original_url_results.json'
    original_urls_path=AUDIT/'original_urls.json'
    results=json.loads(original_results_path.read_text()) if original_results_path.exists() else {}
    valid_original_urls=json.loads(original_urls_path.read_text()) if original_urls_path.exists() else {}
    results={u:r for u,r in results.items() if u in valid_original_urls}
    sync_master(articles,audit_rows,sources,results)
    sync_roadmaps(articles)
    subprocess.run([sys.executable,str(BASE/'generate_posts.py')],check=True)
    print(json.dumps({'articles':len(articles),'sources':len(sources),'confirmed_canonical_mismatches':len(CONFIRMED_MISMATCHES),
                      'original_urls_checked':len(results),'original_http_statuses':dict(Counter(str(r.get('http_status','error')) for r in results.values()))},indent=2))

if __name__=='__main__':main()
