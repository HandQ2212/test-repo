"""Cross-artifact validation; this is not medical certification."""
import csv
from html.parser import HTMLParser
import json
import hashlib
from pathlib import Path
import re
import sys
import openpyxl

BASE=Path(__file__).resolve().parent
ROOT=BASE.parent
sys.path.insert(0,str(BASE))
from rebuild_editorial import CONFIRMED_MISMATCHES
from build_longform import baseline as load_baseline
from image_curation import selections

class Links(HTMLParser):
    def __init__(self):
        super().__init__();self.images=[]
    def handle_starttag(self,tag,attrs):
        if tag=='img':self.images.append(dict(attrs))

def main():
    articles=json.loads((BASE/'articles.json').read_text())
    assert len(articles)==len({a['id'] for a in articles})==83
    js=(BASE/'articles_data.js').read_text().strip()
    assert json.loads(js.removeprefix('const articles = ').removesuffix(';'))==articles
    registry=json.loads((BASE/'data/references.json').read_text())
    sources={s['id']:s for s in registry}
    with (BASE/'data/editorial_sources.tsv').open(newline='') as f:
        expected_sources=list(csv.DictReader(f,delimiter='\t'))
    assert len(sources)==len(registry)==len(expected_sources)==82
    developed=json.loads((BASE/'data/longform_articles.json').read_text())
    assert (BASE/'article_figures.js').is_file()
    assert 'article_figures.js?' in (BASE/'index.html').read_text()
    baseline={a['id']:a for a in load_baseline()}
    selected=selections()
    curated={(s['article_id'],s['role']):s for s in selected}
    catalog={}
    for name in ['editorial_articles.tsv','editorial_women.tsv','editorial_men.tsv']:
        with (BASE/'data'/name).open(newline='') as f:
            for row in csv.DictReader(f,delimiter='\t'):catalog[row['id']]=row
    assert set(catalog)=={a['id'] for a in articles}
    for a in articles:
        aid=a['id'];refs=a['references']
        assert refs and a['source_url']==refs[0]['url']
        assert a['clinical_review_status']=='Not recorded'
        assert a['source_checked_date']=='2026-10-26'
        assert 'not a measured biomarker' in a['metric']
        for r in refs:
            assert r['url']==sources[r['id']]['url'] and r['url'].startswith('https://')
            assert r['support_vi'] and r['support_en'] and r['kind']
        for bodykey,lang in [('content_vi','vi'),('content','en')]:
            body=a[bodykey]
            for field in ['facts','action','limits','care']:assert catalog[aid][field+'_'+lang] in body
            assert body==developed[aid][bodykey]
            assert len(re.findall(r'^## ',body,re.M))>=7
            prose=body.split('## References and what they support')[0].split('## Nguồn tham khảo và phạm vi hỗ trợ')[0]
            assert len(prose.split())>=(600 if lang=='en' else 1000),(aid,lang,'insufficient depth')
            assert '|---' in body.replace(' ',''),(aid,'missing comparison')
            for r in refs:assert r['url'] in body
            assert not re.search(r'TODO|TBD|Lorem ipsum|\[INSERT|example\.com',body)
            assert 'Sơ đồ tóm tắt do biên tập tạo' not in body
            assert 'Original editorial summary' not in body
            assert 'images/editorial/' not in body
        for field in ['hero_image','inline_image']:
            if field=='inline_image' and not a[field]:
                assert not any(a[k] for k in ['inline_caption','inline_caption_vi','inline_image_source_url'])
                continue
            asset=BASE/a[field]
            assert asset.is_file() and asset.suffix.lower() in ['.jpg','.jpeg','.png','.webp']
        assert a['hero_image']==curated.get((aid,'hero'),{}).get('asset',baseline[aid]['hero_image'])
        for role in ['hero','inline']:
            asset=curated.get((aid,role))
            if not asset:continue
            assert a[role+'_image']==asset['asset']
            assert a[role+'_image_source_url']==asset['source_page']
            assert a[role+'_caption']==asset['caption_en'] and a[role+'_caption_vi']==asset['caption_vi']
            assert hashlib.sha256((BASE/asset['asset']).read_bytes()).hexdigest()==asset['sha256']
            assert asset['author'] and asset['license'] and asset['license_url'].startswith('https://')
            assert asset['relevance_vi'] and asset['source_metadata']
        assert a['baseline_revision']=='f68c786'
        page=(BASE/'posts'/f'{aid}.html').read_text();parser=Links();parser.feed(page)
        assert a['title'] in page and a['title_vi'] in page
        assert '../article_figures.js?' in page
        assert len(parser.images)==2*(1+bool(a['inline_image'])),(aid,'unexpected image count')
        assert 'images/editorial/' not in page
        assert 'Sơ đồ tóm tắt do biên tập tạo' not in page
        assert 'Original editorial summary' not in page
        for credit in a['image_attributions']:
            from html import escape
            assert escape(credit['author']) in page
            assert escape(credit['license_url']) in page and escape(credit['source_page']) in page
        for im in parser.images:
            assert im.get('alt'),(aid,'missing alt')
            src=im.get('src','')
            if not src.startswith(('http:','https:','data:')):
                assert (BASE/'posts'/src).resolve().is_file(),(aid,src)
        for r in refs:assert r['url'].replace('&','&amp;') in page or r['url'] in page
        assert 'Clinical Editorial Board' not in page and '100% verified' not in page
        for folder in [BASE/'markdown',ROOT/'Fitness IOS/markdown']:
            if folder.exists():
                md=(folder/f'{aid}.md').read_text()
                assert a['content'] in md and a['content_vi'] in md
    wb=openpyxl.load_workbook(BASE/'docs/Master_Health_Articles_MVP.xlsx')
    ws=wb['Master_Content_MVP'];h={c.value:c.column for c in ws[1]}
    by_id={a['id']:a for a in articles}
    assert ws.max_row==84 and wb['Mapping_Matrix'].max_row==84
    for row in range(2,ws.max_row+1):
        a=by_id[ws.cell(row,h['Content_ID']).value]
        for col,key in [('Title_VI','title_vi'),('Title_EN','title'),('Full_Content_Markdown_VI','content_vi'),('Full_Content_Markdown','content'),('Source_URL','source_url'),('Hero_Image_Asset','hero_image'),('Inline_Figure_Asset','inline_image')]:
            assert (ws.cell(row,h[col]).value or '')==a[key],(a['id'],col)
        if not a['inline_image']:
            for col in ['Inline_Figure_Caption_VI','Inline_Figure_Caption_EN','Inline_Figure_Source_URL']:
                assert not ws.cell(row,h[col]).value and not ws.cell(row,h[col]).hyperlink,(a['id'],col)
        assert ws.cell(row,h['Source_URL']).hyperlink.target==a['source_url']
        assert json.loads(ws.cell(row,h['Reference_List_JSON']).value)==a['references']
        assert json.loads(ws.cell(row,h['Image_Attributions_JSON']).value)==a['image_attributions']
    assert wb['Editorial_Audit'].max_row==84 and wb['Source_Register'].max_row==len(sources)+1
    assert wb['Image_Register'].max_row==len(selected)+1
    for name in ['Master_Health_Articles_MVP.xlsx','Mens_Wellness_Content_Roadmap.xlsx','Master_Health_Articles_MVP - 30 Blog Topics.xlsx']:
        other=ROOT/'Fitness IOS'/name
        if other.exists():assert other.read_bytes()==(BASE/'docs'/name).read_bytes(),name
    before=ROOT/'audit/health-2026-10-26/before/articles.json'
    if before.exists():
        old={a['id']:a for a in json.loads(before.read_text())}
        for aid in CONFIRMED_MISMATCHES:
            assert old[aid]['source_url'] not in by_id[aid]['content']
            assert old[aid]['source_url'] not in by_id[aid]['content_vi']
            assert old[aid]['source_url']!=by_id[aid]['source_url']
    current=ROOT/'audit/health-2026-10-26/revised_url_results.json'
    if current.exists():
        results=json.loads(current.read_text())
        for s in sources.values():
            assert s['url'] in results,s['url']
            assert results[s['url']].get('http_status') in [200,203,403],(s['id'],results[s['url']])
    metadata=ROOT/'audit/health-2026-10-26/revised_pubmed_metadata.json'
    if metadata.exists():
        records={}
        for p in [ROOT/'audit/health-2026-10-26/original_pubmed_metadata.json',metadata]:
            if p.exists():records.update({r['id']:r for r in json.loads(p.read_text())['resultList']['result']})
        for s in sources.values():
            match=re.search(r'pubmed.ncbi.nlm.nih.gov/(\d+)',s['url'])
            if match:
                normalize=lambda t:re.sub(r'[^a-z0-9]','',re.sub(r'\[retracted\]','',t.lower()))
                title=normalize(s['title'].split(' — ',1)[-1])
                assert title in normalize(records[match[1]]['title']),(s['id'],records[match[1]]['title'])
    restored=sum(a['inline_image']==baseline[a['id']]['inline_image'] for a in articles)
    unreviewed=[a for a in articles if not a['image_review_batch']]
    for a in unreviewed:
        assert a['hero_image']==baseline[a['id']]['hero_image']
    assert len(curated)==len(selected)
    print(f'PASS: 83 bilingual articles; {len(sources)} sources; {len(selected)} licensed image selections for {83-len(unreviewed)} articles; {len(unreviewed)} articles pending image sessions; 83 HTML pages; JS/JSON/Markdown/Excel parity; citation checks.')
    print('HTTP access and editorial assessment are separate. Clinical sign-off has not been recorded.')

if __name__=='__main__':main()
